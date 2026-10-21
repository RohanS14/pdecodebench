"""
score_audit.py — compare an expert's grades with the §4.2 judge's on the audit sample.

Reads the expert's grading from either file they may return:
  judge_audit_grading.csv, or judge_audit.xlsx (its "grading" sheet)
and the private key written by select_sample.py. Columns are matched by name.

Refuses to score an incomplete sheet: every item in the key must be graded with a
valid category and both yes/no fields, and a 'some' grade must name the missed
discrepancies.

Reports, pooled (n=20) and per model family (n=10 each):
  agreement on exact category, on "named >=1 real defect" (none vs some/all --
  the paper's claim), on incorrect code claims and on unverifiable execution
  claims: k/n with a 95% Wilson interval, Cohen's kappa (descriptive), confusion
  matrices; the expert's vs the judge's rates on the same items; agreement by the
  expert's confidence; and every disagreement with both explanations.

Usage:
  python frontier/judge_audit/score_audit.py --grading <csv|xlsx> [--key <key.csv>] [--out <dir>]
"""
import argparse
import csv
import math
from collections import Counter
from pathlib import Path

REPO_ROOT = Path(__file__).resolve().parent.parent.parent
PRIVATE = REPO_ROOT / "results" / "frontier" / "judge_audit" / "private"
CATEGORIES = ["none", "some", "all"]
CONFIDENCE = ["low", "medium", "high"]
YES = {"yes", "y", "true", "1"}
NO = {"no", "n", "false", "0"}
FAMILY_LABEL = {"gemini-2.5-flash": "Gemini 2.5 Flash", "gemini-3.7-flash": "Gemini 3.7 Flash"}

# (name, label, how to read the expert's grade, how to read the judge's)
MEASURES = [
    ("category", "Exact category (none / some / all)",
     lambda g: g["category"], lambda k: k["judge_category"]),
    ("named_defect", "Named >=1 real defect (some/all vs none)",
     lambda g: g["category"] != "none", lambda k: k["judge_category"] != "none"),
    ("incorrect_code_claims", "Incorrect code claims",
     lambda g: g["incorrect_code_claims"], lambda k: k["judge_incorrect_code_claims"]),
    ("unverifiable_execution_claims", "Unverifiable execution claims",
     lambda g: g["unverifiable_execution_claims"], lambda k: k["judge_unverifiable_execution_claims"]),
]


# ── input ─────────────────────────────────────────────────────────────────────

def _text(v) -> str:
    return "" if v is None else str(v).strip()


def _yes_no(v):
    t = _text(v).lower()
    return True if t in YES else False if t in NO else None


def read_grading(path: Path) -> list[dict]:
    """Rows of the grading sheet as {column: raw value}, from a .csv or .xlsx."""
    if path.suffix.lower() == ".xlsx":
        from openpyxl import load_workbook
        ws = load_workbook(path, read_only=True, data_only=True)["grading"]
        rows = list(ws.iter_rows(values_only=True))
        header = [_text(h) for h in rows[0]]
        return [dict(zip(header, r)) for r in rows[1:] if any(_text(c) for c in r)]
    with path.open(newline="", encoding="utf-8-sig") as fh:
        return [r for r in csv.DictReader(fh) if any(_text(v) for v in r.values())]


def read_key(path: Path) -> list[dict]:
    with path.open(newline="") as fh:
        key = list(csv.DictReader(fh))
    for k in key:
        for col in ("judge_incorrect_code_claims", "judge_unverifiable_execution_claims"):
            k[col] = _yes_no(k[col])
    return key


def normalize(raw: list[dict], key: list[dict]) -> dict[str, dict]:
    """Validate and normalize the expert's rows; returns {item_id: grade}. Raises
    ValueError listing every problem, so a partly filled sheet is never scored."""
    problems, grades = [], {}
    for r in raw:
        item = _text(r.get("item_id"))
        if not item:
            continue
        if item in grades:
            problems.append(f"{item}: graded more than once")
        g = {
            "category": _text(r.get("category")).lower(),
            "incorrect_code_claims": _yes_no(r.get("incorrect_code_claims")),
            "unverifiable_execution_claims": _yes_no(r.get("unverifiable_execution_claims")),
            "missed_discrepancies": _text(r.get("missed_discrepancies")),
            "explanation": _text(r.get("explanation")),
            "confidence": _text(r.get("confidence")).lower(),
            "minutes_spent": _text(r.get("minutes_spent")),
        }
        if g["category"] not in CATEGORIES:
            problems.append(f"{item}: category must be one of {CATEGORIES}, got {r.get('category')!r}")
        for col in ("incorrect_code_claims", "unverifiable_execution_claims"):
            if g[col] is None:
                problems.append(f"{item}: {col} must be yes/no, got {r.get(col)!r}")
        if g["category"] == "some" and not g["missed_discrepancies"]:
            problems.append(f"{item}: category 'some' requires missed_discrepancies")
        if g["confidence"] and g["confidence"] not in CONFIDENCE:
            problems.append(f"{item}: confidence must be one of {CONFIDENCE}, got {r.get('confidence')!r}")
        grades[item] = g
    expected = {k["item_id"] for k in key}
    if missing := sorted(expected - set(grades)):
        problems.append(f"ungraded items: {', '.join(missing)}")
    if extra := sorted(set(grades) - expected):
        problems.append(f"items not in the key: {', '.join(extra)}")
    if problems:
        raise ValueError("grading sheet is incomplete or invalid:\n  " + "\n  ".join(problems))
    return grades


# ── statistics ────────────────────────────────────────────────────────────────

def wilson(k: int, n: int, z: float = 1.959964) -> tuple[float, float]:
    if n == 0:
        return (float("nan"), float("nan"))
    p = k / n
    centre = (p + z * z / (2 * n)) / (1 + z * z / n)
    half = z * math.sqrt(p * (1 - p) / n + z * z / (4 * n * n)) / (1 + z * z / n)
    return (max(0.0, centre - half), min(1.0, centre + half))


def cohen_kappa(a: list, b: list) -> float | None:
    """None when chance agreement is 1 (both raters used a single identical label),
    where kappa is undefined."""
    n = len(a)
    if n == 0:
        return None
    po = sum(x == y for x, y in zip(a, b)) / n
    ca, cb = Counter(a), Counter(b)
    pe = sum(ca[c] * cb[c] for c in set(ca) | set(cb)) / (n * n)
    return None if pe == 1 else (po - pe) / (1 - pe)


def agreement(pairs: list[tuple]) -> dict:
    k, n = sum(e == j for e, j in pairs), len(pairs)
    lo, hi = wilson(k, n)
    return {"k": k, "n": n, "lo": lo, "hi": hi,
            "kappa": cohen_kappa([e for e, _ in pairs], [j for _, j in pairs])}


def score(grades: dict[str, dict], key: list[dict]) -> dict:
    """All numbers in the report, as plain data (used by the tests)."""
    groups = {"pooled": key}
    for fam in sorted({k["model_family"] for k in key}):
        groups[fam] = [k for k in key if k["model_family"] == fam]
    out = {"groups": {}, "items": []}
    for name, rows in groups.items():
        g = {"n": len(rows), "agreement": {}, "confusion": {}, "rates": {}}
        for m, _, expert, judge in MEASURES:
            pairs = [(expert(grades[k["item_id"]]), judge(k)) for k in rows]
            g["agreement"][m] = agreement(pairs)
            g["confusion"][m] = Counter(pairs)
        for m in ("named_defect", "incorrect_code_claims"):
            _, _, expert, judge = next(x for x in MEASURES if x[0] == m)
            g["rates"][m] = {"expert": sum(bool(expert(grades[k["item_id"]])) for k in rows),
                             "judge": sum(bool(judge(k)) for k in rows), "n": len(rows)}
        out["groups"][name] = g
    by_conf = {}
    for conf in CONFIDENCE + [""]:
        rows = [k for k in key if grades[k["item_id"]]["confidence"] == conf]
        if rows:
            by_conf[conf or "(blank)"] = {
                m: agreement([(expert(grades[k["item_id"]]), judge(k)) for k in rows])
                for m, _, expert, judge in MEASURES[:3]}
    out["by_confidence"] = by_conf
    for k in sorted(key, key=lambda r: r["item_id"]):
        g = grades[k["item_id"]]
        item = {"item_id": k["item_id"], "model_family": k["model_family"], "setting": k["setting"],
                "title": k["title"], "mod_type": k["mod_type"], "expert_confidence": g["confidence"]}
        for m, _, expert, judge in MEASURES:
            item[f"expert_{m}"], item[f"judge_{m}"] = expert(g), judge(k)
            item[f"agree_{m}"] = expert(g) == judge(k)
        item["expert_explanation"], item["expert_missed_discrepancies"] = g["explanation"], g["missed_discrepancies"]
        item["judge_explanation"] = k["judge_explanation"]
        out["items"].append(item)
    return out


# ── report ────────────────────────────────────────────────────────────────────

def _fmt(v) -> str:
    return {True: "yes", False: "no"}.get(v, str(v))


def _pct(k: int, n: int) -> str:
    return f"{k}/{n} ({100 * k / n:.0f}%)" if n else "—"


def render_report(res: dict, grading_path: Path) -> str:
    def group_title(name):
        return "Pooled" if name == "pooled" else FAMILY_LABEL.get(name, name)

    L = ["# Judge audit: expert vs LLM judge", "",
         f"Grading file: `{grading_path.name}`. One expert, {res['groups']['pooled']['n']} items "
         "(5 drawn uniformly at random per thinking setting). Intervals are 95% Wilson; "
         "kappa is descriptive. With n=20 (n=10 per family) intervals are wide: this is a "
         "sanity audit of the judge, not a precise estimate of its accuracy.", ""]

    L += ["## Agreement", "",
          "| measure | " + " | ".join(group_title(g) for g in res["groups"]) + " |",
          "|---|" + "---|" * len(res["groups"])]
    for m, label, _, _ in MEASURES:
        cells = []
        for g in res["groups"].values():
            a = g["agreement"][m]
            kappa = "κ n/a" if a["kappa"] is None else f"κ={a['kappa']:.2f}"
            cells.append(f"{a['k']}/{a['n']} [{100 * a['lo']:.0f}–{100 * a['hi']:.0f}%], {kappa}")
        L.append(f"| {label} | " + " | ".join(cells) + " |")

    L += ["", "## Expert vs judge rates on the same items", "",
          "| | named >=1 real defect: expert | judge | incorrect code claims: expert | judge |",
          "|---|---|---|---|---|"]
    for name, g in res["groups"].items():
        nd, ic = g["rates"]["named_defect"], g["rates"]["incorrect_code_claims"]
        L.append(f"| {group_title(name)} | {_pct(nd['expert'], nd['n'])} | {_pct(nd['judge'], nd['n'])} | "
                 f"{_pct(ic['expert'], ic['n'])} | {_pct(ic['judge'], ic['n'])} |")

    L += ["", "## Confusion matrices (rows: expert, columns: judge)", ""]
    for name, g in res["groups"].items():
        L.append(f"### {group_title(name)}")
        for m, label, _, _ in MEASURES:
            labels = CATEGORIES if m == "category" else [True, False]
            conf = g["confusion"][m]
            L += ["", f"{label}", "", "| expert \\ judge | " + " | ".join(_fmt(x) for x in labels) + " |",
                  "|---|" + "---|" * len(labels)]
            for e in labels:
                L.append(f"| {_fmt(e)} | " + " | ".join(str(conf.get((e, j), 0)) for j in labels) + " |")
        L.append("")

    L += ["## Agreement by expert confidence (pooled)", "",
          "| confidence | category | named >=1 defect | incorrect code claims |", "|---|---|---|---|"]
    for conf, ms in res["by_confidence"].items():
        L.append(f"| {conf} | " + " | ".join(f"{a['k']}/{a['n']}" for a in ms.values()) + " |")

    dis = [it for it in res["items"] if not all(it[f"agree_{m}"] for m, *_ in MEASURES)]
    L += ["", f"## Disagreements ({len(dis)} of {len(res['items'])} items)", ""]
    for it in dis:
        fields = [f"{m}: expert {_fmt(it[f'expert_{m}'])} / judge {_fmt(it[f'judge_{m}'])}"
                  for m, *_ in MEASURES if not it[f"agree_{m}"]]
        L += [f"### {it['item_id']} — {FAMILY_LABEL.get(it['model_family'], it['model_family'])}, "
              f"{it['setting']}, `{it['title']}`", "",
              "- " + "\n- ".join(fields),
              f"- expert confidence: {it['expert_confidence'] or '(blank)'}",
              f"- **expert:** {it['expert_explanation'] or '(none given)'}"
              + (f" Missed: {it['expert_missed_discrepancies']}" if it["expert_missed_discrepancies"] else ""),
              f"- **judge:** {it['judge_explanation']}", ""]
    return "\n".join(L) + "\n"


def write_items_csv(res: dict, path: Path) -> None:
    with path.open("w", newline="") as fh:
        w = csv.DictWriter(fh, fieldnames=list(res["items"][0]))
        w.writeheader()
        w.writerows(res["items"])


def main():
    p = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    p.add_argument("--grading", required=True, type=Path, help="judge_audit_grading.csv or judge_audit.xlsx")
    p.add_argument("--key", type=Path, default=PRIVATE / "judge_audit_key.csv")
    p.add_argument("--out", type=Path, default=REPO_ROOT / "results" / "frontier" / "judge_audit" / "report")
    args = p.parse_args()

    key = read_key(args.key)
    grades = normalize(read_grading(args.grading), key)
    res = score(grades, key)
    args.out.mkdir(parents=True, exist_ok=True)
    stem = args.grading.stem
    (args.out / f"{stem}_report.md").write_text(render_report(res, args.grading))
    write_items_csv(res, args.out / f"{stem}_items.csv")
    pooled = res["groups"]["pooled"]["agreement"]
    print("pooled agreement: " + ", ".join(f"{m} {a['k']}/{a['n']}" for m, a in pooled.items()))
    print(f"wrote {args.out / f'{stem}_report.md'} and {args.out / f'{stem}_items.csv'}")


if __name__ == "__main__":
    main()
