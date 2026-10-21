"""
build_packet.py — render the human-audit materials from the private key.

Every block the expert sees is built by the judge's own helpers in
frontier/judge_invalidity_reasoning.py from the same dataset file and the same
Stage 2 justification the judge was given, so the expert and the judge see the
same inputs. tests/test_judge_audit.py checks each block appears verbatim in
build_judge_prompt()'s output. Nothing identifying is shown: no model, thinking
setting, row title or judge output.

Writes results/frontier/judge_audit/share/ (sent to the expert):
  judge_audit_packet.html   instructions + rubric + 20 item cards
  judge_audit_grading.csv   one row per item to fill in
  judge_audit.xlsx          instructions / grading (dropdowns) / items (content) sheets

Usage:
  python frontier/judge_audit/build_packet.py
"""
import csv
import html
import json
import sys
from pathlib import Path

REPO_ROOT = Path(__file__).resolve().parent.parent.parent
sys.path.insert(0, str(REPO_ROOT))
sys.path.insert(0, str(REPO_ROOT / "frontier"))

from frontier.judge_invalidity_reasoning import (  # noqa: E402
    JUDGE_DECL,
    build_caveat_note,
    find_ground_truth_reference,
    find_valid_counterpart_code,
)
from frontier.judge_audit.select_sample import PRIVATE, OUT_ROOT, dataset_csv  # noqa: E402
from shared.dataset_io import load_dataset  # noqa: E402
from upload_traces import arm_files  # noqa: E402

SHARE = OUT_ROOT / "share"
PDE_NAME = {"burgers": "Burgers", "heat": "Heat", "wave": "Wave", "navier-stokes": "Navier–Stokes"}
GRADING_COLUMNS = ["item_id", "pde", "category", "incorrect_code_claims", "unverifiable_execution_claims",
                   "missed_discrepancies", "explanation", "confidence", "minutes_spent"]
CONTENT_COLUMNS = ["caveat_note", "valid_code", "corrupted_code", "reference_valid_code",
                   "reference_corrupted_code", "invalidity_note", "model_justification"]
CHOICES = {"category": ["none", "some", "all"], "incorrect_code_claims": ["yes", "no"],
           "unverifiable_execution_claims": ["yes", "no"], "confidence": ["low", "medium", "high"]}

_PROPS = JUDGE_DECL["parameters_json_schema"]["properties"]
RUBRIC = [  # the judge's own field definitions, verbatim
    ("category", "none / some / all", _PROPS["category"]["description"]),
    ("incorrect_code_claims", "yes / no", _PROPS["contains_incorrect_code_claims"]["description"]),
    ("unverifiable_execution_claims", "yes / no", _PROPS["contains_unverifiable_execution_claims"]["description"]),
    ("missed_discrepancies", "free text", "Required when category is 'some': name the real discrepancies "
     "between the valid and corrupted code that the justification did not mention."),
    ("explanation", "free text", "1-3 sentences justifying your three choices above."),
    ("confidence", "low / medium / high", "How confident you are in your grading of this item."),
    ("minutes_spent", "number, optional", "Approximate minutes spent on this item."),
]
INSTRUCTIONS = [
    "Each item shows a PDE solver that has been deliberately corrupted so that it no longer "
    "implements valid physics or numerics, next to the valid version of the same solver. A "
    "language model examined the corrupted code and correctly answered that it is invalid. "
    "Your task is to judge whether the model's stated justification identifies the real "
    "difference(s) between the valid and corrupted code.",
    "Compare the valid and corrupted code directly to determine what actually changed; no diff "
    "is provided. Where the corrupted code has no comments, obfuscated variable names, or "
    "misleading comments, a note says so and readable reference versions (original names and "
    "genuine comments) are shown below it.",
    "The annotator's invalidity note describes the general symptom the corruption is expected to "
    "cause. It is a general description, not necessarily the specific mechanism -- do not treat "
    "it as the definitive or only correct answer.",
    "The model was able to execute the code before giving its answer, so claims about runtime "
    "behaviour may be genuine observations rather than speculation. You cannot verify them -- you "
    "have the code, not the execution output. Judge claims about the code on their correctness; "
    "record unverifiable runtime figures in the separate field rather than counting them as errors.",
    "The model was asked for a binary answer and a validity score, and its justification may "
    "mention the score. Evaluate the claims in the justification solely with respect to the actual "
    "invalidities in the code, not the stated score.",
    "Record your grades in judge_audit_grading.csv (or the 'grading' sheet of judge_audit.xlsx), "
    "one row per item ID. Please grade each item independently and in any order.",
]


def build_items() -> list[dict]:
    """One dict per key row: item_id, pde, and the judge's inputs for that row."""
    with (PRIVATE / "judge_audit_key.csv").open() as fh:
        key = list(csv.DictReader(fh))
    datasets, episodes, items = {}, {}, []
    for k in key:
        s = k["setting"]
        if s not in datasets:
            datasets[s] = load_dataset(str(dataset_csv(s)))
            with arm_files(s)["episodes"].open() as fh:
                episodes[s] = {r["title"]: r for r in map(json.loads, fh)}
        df, ep = datasets[s], episodes[s][k["title"]]
        row = df[df["title"] == k["title"]].iloc[0]
        valid_code = find_valid_counterpart_code(df, row["gt_sample"], row["mod_type"])
        ref_valid, ref_invalid = find_ground_truth_reference(df, row["gt_sample"], row["mod_type"])
        items.append({
            "item_id": k["item_id"],
            "pde": PDE_NAME.get(str(row["pde_class"]).lower(), str(row["pde_class"])),
            "caveat_note": build_caveat_note(row["mod_type"], row.get("corruption_source_pde")),
            "valid_code": valid_code,
            "corrupted_code": str(row["code"]),
            "reference_valid_code": ref_valid or "",
            "reference_corrupted_code": ref_invalid or "",
            "invalidity_note": str(row.get("invalidity_note") or ""),
            "model_justification": str(ep.get("s2_submit_args", {}).get("valid_exp") or ""),
        })
        assert items[-1]["valid_code"] and items[-1]["model_justification"], k["item_id"]
    return items


# ── HTML ──────────────────────────────────────────────────────────────────────

CSS = """
:root { --bg:#fbfaf7; --fg:#1d1d1b; --muted:#5f5e58; --card:#ffffff; --line:#e3e0d8; --code:#f4f2ec;
        --accent:#8a4b16; --note:#fff6e0; }
@media (prefers-color-scheme: dark) { :root { --bg:#18181a; --fg:#ececea; --muted:#a3a29c; --card:#212124;
        --line:#36363a; --code:#1b1b1e; --accent:#e0a36b; --note:#2e2817; } }
* { box-sizing:border-box; }
body { margin:0; background:var(--bg); color:var(--fg); font:16px/1.55 -apple-system, "Segoe UI", sans-serif; }
main { max-width:1500px; margin:0 auto; padding:24px 16px 80px; }
h1 { font-size:1.6rem; margin:0 0 4px; } h2 { font-size:1.25rem; margin:0; }
h3 { font-size:.85rem; text-transform:uppercase; letter-spacing:.04em; color:var(--muted); margin:20px 0 6px; }
.intro, .card { background:var(--card); border:1px solid var(--line); border-radius:10px; padding:20px; margin:18px 0; }
.rubric { border-collapse:collapse; width:100%; } .rubric td { border-top:1px solid var(--line); padding:8px; vertical-align:top; }
.rubric td:first-child { font-family:ui-monospace, Menlo, monospace; white-space:nowrap; }
.toc a { margin-right:10px; color:var(--accent); }
.card header { display:flex; gap:12px; align-items:baseline; border-bottom:1px solid var(--line); padding-bottom:10px; }
.pde { color:var(--muted); }
.note { background:var(--note); border-radius:6px; padding:10px 12px; margin:12px 0; }
.pair { display:grid; grid-template-columns:1fr 1fr; gap:12px; }
@media (max-width:900px) { .pair { grid-template-columns:1fr; } }
pre { background:var(--code); border:1px solid var(--line); border-radius:6px; margin:0; padding:10px 0;
      overflow-x:auto; font:12.5px/1.45 ui-monospace, Menlo, monospace; max-height:70vh; }
pre code { display:block; counter-reset:ln; }
pre code span { display:block; padding:0 12px 0 0; white-space:pre; }
pre code span::before { counter-increment:ln; content:counter(ln); display:inline-block; width:3.2em;
      padding-right:.8em; text-align:right; color:var(--muted); user-select:none; }
.justification { white-space:pre-wrap; border-left:3px solid var(--accent); padding:6px 12px; }
"""


def code_block(code: str) -> str:
    lines = "".join(f"<span>{html.escape(line) or ' '}</span>" for line in code.split("\n"))
    return f"<pre><code>{lines}</code></pre>"


def card(it: dict) -> str:
    parts = [f'<section class="card" id="{it["item_id"]}"><header><h2>{it["item_id"]}</h2>'
             f'<span class="pde">{html.escape(it["pde"])}</span></header>']
    if it["caveat_note"]:
        parts.append(f'<div class="note">{html.escape(it["caveat_note"])}</div>')
    parts.append('<div class="pair"><div><h3>Valid code (base)</h3>' + code_block(it["valid_code"]) +
                 '</div><div><h3>Corrupted code (what the model evaluated)</h3>' +
                 code_block(it["corrupted_code"]) + "</div></div>")
    if it["reference_valid_code"] or it["reference_corrupted_code"]:
        parts.append('<div class="pair"><div><h3>Readable reference: valid</h3>' +
                     code_block(it["reference_valid_code"]) +
                     '</div><div><h3>Readable reference: corrupted</h3>' +
                     code_block(it["reference_corrupted_code"]) + "</div></div>")
    parts.append(f'<h3>Annotator\'s invalidity note (general symptom)</h3><div class="note">'
                 f'{html.escape(it["invalidity_note"])}</div>')
    parts.append(f'<h3>Model\'s justification</h3><div class="justification">'
                 f'{html.escape(it["model_justification"])}</div></section>')
    return "".join(parts)


def render_html(items: list[dict]) -> str:
    rubric = "".join(f"<tr><td>{c}</td><td>{html.escape(v)}</td><td>{html.escape(d)}</td></tr>"
                     for c, v, d in RUBRIC)
    toc = "".join(f'<a href="#{it["item_id"]}">{it["item_id"]}</a>' for it in items)
    intro = "".join(f"<p>{html.escape(p)}</p>" for p in INSTRUCTIONS)
    return (f'<!doctype html><html lang="en"><head><meta charset="utf-8">'
            f'<meta name="viewport" content="width=device-width, initial-scale=1">'
            f"<title>Judge Audit Packet</title><style>{CSS}</style></head><body><main>"
            f"<h1>Justification audit: {len(items)} items</h1>"
            f'<section class="intro"><h2>Instructions</h2>{intro}'
            f'<h3>Grading fields</h3><table class="rubric">{rubric}</table>'
            f'<h3>Items</h3><div class="toc">{toc}</div></section>'
            + "".join(card(it) for it in items) + "</main></body></html>")


# ── CSV / xlsx ────────────────────────────────────────────────────────────────

def write_grading_csv(items: list[dict], path: Path) -> None:
    with path.open("w", newline="") as fh:
        w = csv.DictWriter(fh, fieldnames=GRADING_COLUMNS)
        w.writeheader()
        for it in items:
            w.writerow({"item_id": it["item_id"], "pde": it["pde"]})


def write_xlsx(items: list[dict], path: Path) -> None:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font
    from openpyxl.worksheet.datavalidation import DataValidation

    wb = Workbook()
    ins = wb.active
    ins.title = "instructions"
    ins.column_dimensions["A"].width = 30
    ins.column_dimensions["B"].width = 18
    ins.column_dimensions["C"].width = 110
    r = 1
    for p in INSTRUCTIONS:
        ins.cell(r, 1, p).alignment = Alignment(wrap_text=True, vertical="top")
        ins.merge_cells(start_row=r, start_column=1, end_row=r, end_column=3)
        ins.row_dimensions[r].height = 60
        r += 1
    r += 1
    for c, h in enumerate(["field", "values", "definition"], 1):
        ins.cell(r, c, h).font = Font(bold=True)
    for field, values, desc in RUBRIC:
        r += 1
        for c, v in enumerate([field, values, desc], 1):
            ins.cell(r, c, v).alignment = Alignment(wrap_text=True, vertical="top")

    g = wb.create_sheet("grading")
    g.append(GRADING_COLUMNS)
    for it in items:
        g.append([it["item_id"], it["pde"]] + [None] * (len(GRADING_COLUMNS) - 2))
    n = len(items) + 1
    for col, choices in CHOICES.items():
        letter = chr(ord("A") + GRADING_COLUMNS.index(col))
        dv = DataValidation(type="list", formula1='"' + ",".join(choices) + '"', allow_blank=True)
        g.add_data_validation(dv)
        dv.add(f"{letter}2:{letter}{n}")
    for col, width in zip("ABCDEFGHI", [8, 14, 10, 14, 16, 40, 50, 11, 10]):
        g.column_dimensions[col].width = width
    for row in g.iter_rows(min_row=1, max_row=n):
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical="top")
    for cell in g[1]:
        cell.font = Font(bold=True)
    g.freeze_panes = "C2"

    s = wb.create_sheet("items")
    s.append(["item_id", "pde"] + CONTENT_COLUMNS)
    for it in items:
        s.append([it["item_id"], it["pde"]] + [it[c] for c in CONTENT_COLUMNS])
    for col in "CDEFGHI":
        s.column_dimensions[col].width = 60
    for row in s.iter_rows():
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical="top")
            if cell.column_letter in "DEFG":
                cell.font = Font(name="Menlo", size=9)
    for cell in s[1]:
        cell.font = Font(bold=True)
    s.freeze_panes = "C2"
    s.protection.sheet = True  # read-only content; no password, just guards against accidental edits
    wb.active = 1
    wb.save(path)


def main():
    items = build_items()
    SHARE.mkdir(parents=True, exist_ok=True)
    (SHARE / "judge_audit_packet.html").write_text(render_html(items))
    write_grading_csv(items, SHARE / "judge_audit_grading.csv")
    write_xlsx(items, SHARE / "judge_audit.xlsx")
    print(f"wrote {len(items)} items to {SHARE}/: judge_audit_packet.html, judge_audit_grading.csv, judge_audit.xlsx")


if __name__ == "__main__":
    main()
